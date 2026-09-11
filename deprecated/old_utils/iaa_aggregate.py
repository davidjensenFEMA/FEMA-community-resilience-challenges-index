# -*- coding: utf-8 -*-
"""
Created on Tue Jun  1 07:48:42 2021

@author: jhutchison

# https://realpython.com/openpyxl-excel-spreadsheets-python/#adding-formulas
conda install -c anaconda xlwings # comes with anaconda/unnecessary check
conda install -c anaconda openpyxl
# openpyxl allows writing formulas/formatting to output

"""

# %% Packages

import glob

# import matplotlib.pyplot as plt
import matplotlib.colors
import numpy as np
import pandas as pd

# import pickle
import re

# import seaborn as sb
import time
import xlwings as xw

from openpyxl import Workbook  # load_workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from tqdm import tqdm

# for categorization
# from nltk.corpus import stopwords
# from nltk.stem.snowball import SnowballStemmer
# from sklearn.cluster import KMeans

# for categorization method evaluation
# from sklearn.model_selection import RepeatedKFold
# from sklearn.neighbors import KNeighborsClassifier
# from sklearn.feature_extraction.text import TfidfVectorizer
# from sklearn.metrics import accuracy_score
# from skmultilearn.problem_transform import BinaryRelevance

##############################################################################
# %% Functions


def excel_columns(columns):
    col_idx = {}
    alphabet = list(map(chr, range(65, 91)))
    for idx, col in enumerate(columns):
        # workbook will write with an index column as column zero (column A)
        # so permanently shift col idx by 1
        idx += 1
        if idx < len(alphabet):
            col_idx[col] = alphabet[idx]
        elif idx < len(alphabet) * (len(alphabet) + 1):
            # method should work up to ZZ in excel (702 columns)
            ref = np.floor(idx / len(alphabet))
            # format ref as int with 0 index
            ref = int(ref) - 1
            idx = idx % len(alphabet)
            col_idx[col] = alphabet[int(ref)] + alphabet[idx]
        else:
            print("too many columns")
    return col_idx


##############################################################################
# %% Global Variables

time_start = time.time()

path_data = "data/" + r"Values/"
# path_data = 'data/' + r'Rollups/'
path_out = "outputs/"

# Elicitation topic, this is the sheet name for the workbook collection
# folder_name = 'Security Officer Profile'
# sheet_name = 'Security Officer Profile'
# sheet_range = 'A1:L229'
# template_file = 'V6 Relative Importance Elicitation - Security Officer' +\
#     ' Profile_2021-05-26.xlsx'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

# folder_name = 'Security Management Profile'
# sheet_name = 'Security Management Profile'
# sheet_range = 'A1:L443'
# template_file = 'V6 Relative Importance Elicitation - Security Management' +\
#     ' Profile_2021-06-04'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

folder_name = "Physical Security - Fencing"
sheet_name = "Physical Security - Fencing"
sheet_range = "A1:L130"
template_file = (
    "V6 Relative Importance Elicitation - Physical Security" + " - Fencing_2021-06-10_3"
)
survey_columns = {
    "V6 Rank": "Rank",
    "V6 Relative Importance Range": "Importance",
    "Relative Importance Selection Confidence Level"
    + "\n*Definitions at end of page": "Confidence",
    "PSA Comments": "Comments",
}

# folder_name = 'Physical Security - Gates'
# sheet_name = 'Physical Security - Gates'
# sheet_range = 'A1:L110'
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - Gates_2021-06-22'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

# folder_name = 'Physical Security - Parking'
# sheet_name = 'PhysicalSecurity-ParkingBarrier'
# sheet_range = 'A1:L217'
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - Parking and Barriers_2021-07-07_2'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

# folder_name = 'Physical Security - Building Envelope'
# sheet_name = 'PhysicalSecurity-BuildingEnvelo'
# sheet_range = 'A1:L247'
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - Building Envelope_2021-07-28_2.xlsx'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }


# folder_name = 'Physical Security - VSS'
# sheet_name = 'Physical Security-VSS'
# sheet_range = 'A1:L69'
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - VSS_2021-08-12.xlsx'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

# folder_name = 'Physical Security - IDS'
# sheet_name = 'Physical Security-IDS'
# sheet_range = 'A1:L155'
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - IDS_2021-08-12_2.xlsx'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

# folder_name = 'Physical Security - Illumination'
# sheet_name = 'Physical Security-Illumination'
# sheet_range = 'A1:L176'
# print('wrong template file for illumination')
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - IIlumination_2021-08-12_3.xlsx'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

##############################################################################
# folder_name = 'Physical Security - Fencing'
# sheet_name = 'Physical Security - Fencing'
# sheet_range = 'A1:L157'
# template_file = 'V6 Rollup Elicitation - Physical Security' +\
#     ' - Fencing_2021-09-09.xlsx'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

# folder_name = 'Physical Security - Gates'
# sheet_name = 'Physical Security - Gates'
# sheet_range = 'A1:L150'
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - Gates_2021-09-17'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

# folder_name = 'Physical Security - Illumination'
# sheet_name = 'Physical Security-Illumination'
# sheet_range = 'A1:L291'
# template_file = 'V6 Relative Importance Elicitation - Physical Security' +\
#     ' - Illumination_2021-09-20.xlsx'
# survey_columns = {
#     'V6 Rank': 'Rank',
#     'V6 Relative Importance Range': 'Importance',
#     'Relative Importance Selection Confidence Level' +
#     '\n*Definitions at end of page': 'Confidence',
#     'PSA Comments': 'Comments'
#     }

##############################################################################
# %% Local Variables

# Build replacement dictionaries for string responses

# Relative importance
# Build dictionary of labels to replace {'6-10': 8}
ranges = {}
for num in range(1, 96, 5):
    key = str(num) + "-" + str(num + 4)
    ranges[key] = num + 2.0
num = 96
key = str(num) + "-" + str(num + 3)
ranges[key] = num + 1.5

# Confidence level
levels = {
    "Fully Confident": 4,
    "Mostly Confident": 3,
    "Partially Confident": 2,
    "Slightly Confident": 1,
    "Not Confident": 0,
}

# Consistent Format ('_' preferred by python, ' ' in excel)
sep = " "

# Conditional Formatting Thresholds
threshold_importance = 20  # std
threshold_confidence = 2.5  # mean

##############################################################################
# %% Read xlsx

# List all files
files_xlsx = glob.glob(path_data + folder_name + r"/*.xlsx", recursive=False)
# In the future, sort these files alphabetically!
files_xlsx = [file for file in files_xlsx if r"\~$" not in file]

# Open each workbook (they must be closed to run this script)
# Save results with the file identifier (last '_') as 'records'
records = []
for file_name in tqdm(files_xlsx):
    # uniform labeling will be helpful, this method won't universally apply
    file_id = file_name.split("_", -1)[-1][:-5]
    file_id = re.sub(" ", sep, file_id)  # inactive if sep == ' '
    wb = xw.Book(file_name, password="DHSotof")
    sheet = wb.sheets[sheet_name]
    df = sheet[sheet_range].options(pd.DataFrame, index=False, header=True).value
    df.index.name = file_id
    records.append(df)
    # Close workbook if more then one workbook is open.
    # you won't get empty grey excel app since you have another workbook open.
    # ** Update ** This doesn't work, it's going to quit every time **
    excel_app = xw.apps.active
    if xw.apps.count > 1:
        wb.close()
    # close excel application if only one workbook is open
    else:
        excel_app.quit()

##############################################################################
# %% Aggregate records

# For each record, collect the responses, column subset = columns_survey
responses = pd.Series(dtype="object", index=survey_columns.values())
for idx, record in enumerate(records):
    df = record[survey_columns.keys()]
    for survey, survey_abb in survey_columns.items():
        # Initialize dataframe for each type of response on first record
        if idx == 0:
            responses[survey_abb] = pd.DataFrame()
        responses[survey_abb][df.index.name] = df[survey]
        responses[survey_abb].index.name = survey_abb
        # Replace string labels with floats
        if survey_abb == "Importance":
            ser = df[survey].copy(deep=True)
            ser = ser.replace(ranges, regex=False)
            ser = pd.to_numeric(ser, errors="coerce")
            responses[survey_abb][df.index.name] = ser
        elif survey_abb == "Confidence":
            ser = df[survey].copy(deep=True)
            ser = ser.replace(levels, regex=False)
            ser = pd.to_numeric(ser, errors="coerce")
            responses[survey_abb][df.index.name] = ser

###############################################################################
# %% Analysis
# Current: Mean and Median by Survey
# Future: NLP on comments!

# create a blank template
blank = records[0].copy(deep=True)
# drop elicted response from template
blank = blank.drop(survey_columns.keys(), axis=1)


# df = responses[list(survey_columns.values())[1]]

# Modify response dataframe so column names are legible
for survey_abb in survey_columns.values():
    df = responses[survey_abb]
    cols = list(df.columns)
    cols = [(survey_abb + sep + col) for col in cols]
    df.columns = cols
    # Calculate Output Statistics
    if survey_abb != "Comments":
        col_mean = survey_abb + sep + "Mean"
        col_std = survey_abb + sep + "Std Dev"
        # df = df[df.columns].apply(pd.to_numeric, errors='coerce'). \
        # fillna(0).astype(float).dropna()
        df[col_mean] = df.mean(axis=1)
        df[col_std] = df.std(axis=1)
        # df[(survey_abb + sep + 'Median')] = df.median(axis=1)
        # df[(survey_abb + sep + 'Mode')] = df.mode(axis=1)
    # Reformat Output Statistics as Excel Columns
    if survey_abb != "Comments":
        columns = [blank.columns, " ", cols]
        columns = [item for sublist in columns for item in sublist]
        col_idx = excel_columns(columns)
        col_range = col_idx[cols[0]] + ":" + col_idx[cols[-1]]
        rows = df[col_mean].dropna().index
        for row in rows:
            # output includes header row and 1 index, add 2 to rows to adjust
            ref = row + 2
            array_start = col_idx[cols[0]] + str(ref)
            array_end = col_idx[cols[-1]] + str(ref)
            form = "AVERAGE"
            cell_val = form + "(" + array_start + ":" + array_end + ")"
            df.loc[row, col_mean] = "=ROUND(" + cell_val + ", 2)"
            form = "_xlfn.STDEV.P"
            cell_val = form + "(" + array_start + ":" + array_end + ")"
            df.loc[row, col_std] = "=ROUND(" + cell_val + ", 2)"
    # insert empty/offset column for readability
    empty_col = pd.Series(index=df.index, dtype="object")
    empty_col.name = " "
    blank = pd.concat([blank, empty_col, df], axis=1)

###############################################################################
# %% Conditional Formatting
# openpyxl formula and formatting
"""
workbook = Workbook()
workbook.sheetnames
sheet = workbook.active
sheet.title = sheet_name
cell = sheet["A1"]
cell.value = "hey"
cell.value

# You can also define the position to create the sheet at
hr_sheet = workbook.create_sheet("HR", 0)

# To remove them, just pass the sheet as an argument to the .remove()

workbook.remove(hr_sheet)
workbook.sheetnames

# freeze panes
sheet.freeze_panes = "C2"
"""
# Build the workbook
wb = Workbook()
ws = wb.active
df = blank.copy(deep=True)
df.shape

for idx, r in tqdm(enumerate(dataframe_to_rows(df, index=True, header=True))):
    # function inserts empty row after header, skip that empty row
    if idx != 1:
        ws.append(r)

# colors
colors_rgb = {
    "orange_spec": matplotlib.colors.to_hex((255 / 255, 113 / 255, 40 / 255)),
    "gold": matplotlib.colors.to_hex((255 / 255, 230 / 255, 153 / 255)),
    "red": matplotlib.colors.to_hex((248 / 255, 105 / 255, 107 / 255)),
    "blue": matplotlib.colors.to_hex((90 / 255, 138 / 255, 198 / 255)),
    "white": matplotlib.colors.to_hex((252 / 255, 252 / 255, 255 / 255)),
    "yellow": matplotlib.colors.to_hex((255 / 255, 255 / 255, 0 / 255)),
    "orange": matplotlib.colors.to_hex((255 / 255, 192 / 255, 0 / 255)),
    "orange_acc2": matplotlib.colors.to_hex((237 / 255, 129 / 255, 49 / 255)),
}

colors_rgb = {key: val.replace("#", "") for key, val in colors_rgb.items()}

# Apply formatting

color_scale_mean = ColorScaleRule(
    start_type="min",
    start_color=colors_rgb["orange_spec"],
    end_type="max",
    end_color=colors_rgb["gold"],
)
color_scale_std = ColorScaleRule(
    start_type="min",
    start_color=colors_rgb["gold"],
    end_type="max",
    end_color=colors_rgb["orange_spec"],
)
color_scale_imp = ColorScaleRule(
    start_type="num",
    start_value=0,
    start_color=colors_rgb["red"],
    mid_type="num",
    mid_value=50,
    mid_color=colors_rgb["white"],
    end_type="num",
    end_value=100,
    end_color=colors_rgb["blue"],
)
color_scale_con = ColorScaleRule(
    start_type="num",
    start_value=0,
    start_color=colors_rgb["red"],
    mid_type="num",
    mid_value=2,
    mid_color=colors_rgb["white"],
    end_type="num",
    end_value=4,
    end_color=colors_rgb["blue"],
)

rows = df[col_mean].dropna().index
col_idx = excel_columns(list(df.columns))

# Importance Agreement (standard deviation)
col_excel = col_idx["Importance Std Dev"]
array_start = col_excel + str(rows[0] + 2)
array_end = col_excel + str(rows[-1] + 2)
array_text = array_start + ":" + array_end
ws.conditional_formatting.add(array_text, color_scale_std)

# Confidence Level (mean)
col_excel = col_idx["Confidence Mean"]
array_start = col_excel + str(rows[0] + 2)
array_end = col_excel + str(rows[-1] + 2)
array_text = array_start + ":" + array_end
ws.conditional_formatting.add(array_text, color_scale_mean)

# Importance Values
cols = responses["Importance"].columns
array_start = col_idx[cols[0]] + str(rows[0] + 2)
array_end = col_idx[cols[-3]] + str(rows[-1] + 2)
array_text = array_start + ":" + array_end
ws.conditional_formatting.add(array_text, color_scale_imp)

# Confidence Values
cols = responses["Confidence"].columns
array_start = col_idx[cols[0]] + str(rows[0] + 2)
array_end = col_idx[cols[-3]] + str(rows[-1] + 2)
array_text = array_start + ":" + array_end
ws.conditional_formatting.add(array_text, color_scale_con)

###############################################################################
# %% Analysis

# Trouble (rule)
# Create fill
fill_orange_drk = PatternFill(
    start_color=colors_rgb["orange_acc2"],
    end_color=colors_rgb["orange_acc2"],
    fill_type="solid",
)

col_excel = col_idx["Importance Std Dev"]
array_start = col_excel + str(rows[0] + 2)
formula_text1 = array_start + ">" + str(threshold_importance)

col_excel = col_idx["Confidence Mean"]
array_start = col_excel + str(rows[0] + 2)
formula_text2 = array_start + "<" + str(threshold_confidence)
formula_text = "AND(" + formula_text1 + "," + formula_text2 + ")"
formula_rule_confidence = FormulaRule(formula=[formula_text], fill=fill_orange_drk)


# Trouble (application)
# col_excel = col_idx['Questions/Answers']
col_excel = col_idx[" "]
array_start = col_excel + str(rows[0] + 2)
array_end = col_excel + str(rows[-1] + 2)
array_text = array_start + ":" + array_end
ws.conditional_formatting.add(array_text, formula_rule_confidence)


# Importance Agreement (rule)
# Create fill
fill_yellow = PatternFill(
    start_color=colors_rgb["yellow"], end_color=colors_rgb["yellow"], fill_type="solid"
)

col_excel = col_idx["Importance Std Dev"]
array_start = col_excel + str(rows[0] + 2)
formula_text = array_start + ">" + str(threshold_importance)
formula_rule_agreement = FormulaRule(formula=[formula_text], fill=fill_yellow)


# Importance Agreement (application)
# col_excel = col_idx['Questions/Answers']
col_excel = col_idx[" "]
array_start = col_excel + str(rows[0] + 2)
array_end = col_excel + str(rows[-1] + 2)
array_text = array_start + ":" + array_end
ws.conditional_formatting.add(array_text, formula_rule_agreement)


# Confidence (rule)
# Create fill
fill_orange = PatternFill(
    start_color=colors_rgb["orange"], end_color=colors_rgb["orange"], fill_type="solid"
)

col_excel = col_idx["Confidence Mean"]
array_start = col_excel + str(rows[0] + 2)
formula_text1 = array_start + "<" + str(threshold_confidence)
formula_text2 = array_start + ">0"
formula_text = "AND(" + formula_text1 + "," + formula_text2 + ")"
formula_rule_confidence = FormulaRule(formula=[formula_text], fill=fill_orange)


# Confidence (application)
# col_excel = col_idx['Questions/Answers']
col_excel = col_idx[" "]
array_start = col_excel + str(rows[0] + 2)
array_end = col_excel + str(rows[-1] + 2)
array_text = array_start + ":" + array_end
ws.conditional_formatting.add(array_text, formula_rule_confidence)

###############################################################################
# %% Write

file_out = sheet_name.lower().replace(" ", "_") + ".xlsx"
wb.save(filename=path_out + "rollup_" + file_out)


###############################################################################
# %% Time

time_end = time.time()
print("Total time %s" % (time_end - time_start))
